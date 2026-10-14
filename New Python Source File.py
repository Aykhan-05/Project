import os
from openpyxl import Workbook, load_workbook
import tkinter as tk
from tkinter import ttk

filepath = "All_Group_Anbar_Satis_Sistemi_v1.xlsx"

def open_excel():
    # workbook = load_workbook(filepath)

    # print(workbook.sheetnames)
    os.startfile(filepath)

window = tk.Tk()
window.title("Storage Management")
window.geometry("900x600")

title = ttk.Label(
    window,
    text="Storage Management",
    font=("Arial", 20)
)

title.pack(pady=20)

add_button = ttk.Button(
    window,
    text="Open Excel",
    command=open_excel
)
add_button.pack(pady=10)

stock_button = ttk.Button(
    window,
    text="Current Stock"
)
stock_button.pack(pady=10)

window.mainloop()


